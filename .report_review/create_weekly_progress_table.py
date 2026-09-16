from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "docs" / "Weekly_Progress_Report.docx"


HEADERS = ["Week", "Date", "Development Progress", "Person(s) Responsible", "Status"]

ROWS = [
    [
        "1",
        "15-21 July 2026",
        "The first dungeon map was created. The initial environment layout and basic project structure were prepared to provide a foundation for the game.",
        "Huang Wan Jun",
        "Completed",
    ],
    [
        "2",
        "22-28 July 2026",
        "The player character and eight-directional movement were implemented together with the camera system. The first enemy prototype, Knife weapon, invisible walls and basic map boundaries were also added.",
        "Ng Poh Hui, Tiu Han Xuen, Angel Yap Yoon Ying and Huang Wan Jun",
        "Completed",
    ],
    [
        "3",
        "29 July-4 August 2026",
        "The map boundary script was completed to prevent the player from leaving the playable area. The Gun weapon was implemented, and the completed features were integrated for the first project milestone.",
        "Angel Yap Yoon Ying and Huang Wan Jun",
        "Completed",
    ],
    [
        "4",
        "5-11 August 2026",
        "The enemy spawning section for Level 1 was developed. Enemy placement and wave-related functions were prepared to support the survival gameplay system.",
        "Huang Wan Jun and Tiu Han Xuen",
        "Completed",
    ],
    [
        "5",
        "12-18 August 2026",
        "The user interface, experience orb logic and player scripts were improved. The Level Up interface and enemy reward balance were refined. Level 2, updated enemies and the Skeleton enemy were also added.",
        "All team members",
        "Completed",
    ],
    [
        "6",
        "19-25 August 2026",
        "The ranged enemy was implemented to introduce long-distance enemy attacks. The enemy, player and weapon systems were merged and tested together, and compatibility issues between the different systems were corrected.",
        "Tiu Han Xuen and Angel Yap Yoon Ying",
        "Completed",
    ],
    [
        "7",
        "26 August-1 September 2026",
        "The weapon system was expanded with additional weapons and weapon behaviour changes. The Weapon Store was introduced to support weapon unlocking and upgrading. Weapon 6 was added, while unnecessary files were removed from the project.",
        "Angel Yap Yoon Ying",
        "Completed",
    ],
    [
        "8",
        "2-8 September 2026",
        "Level 3 traps were developed, including spike rows, fire sweeps and falling rocks. Safe spawning, the minimap, boss arena trial, Level 3 and the boss level were implemented. The menu, player profile, HUD, audio feedback and Final Boss were also introduced.",
        "All team members",
        "Completed",
    ],
    [
        "9",
        "9-12 September 2026",
        "Weapon behaviour was updated, additional weapons were completed, and weapon-related bugs were fixed. Final Boss issues, enemy bugs and the Level 1 bomb problem were also corrected during final gameplay testing.",
        "All team members",
        "Completed",
    ],
    [
        "10",
        "13-16 September 2026",
        "The project files were reorganized and the technical report was prepared. Code explanations, screenshots, figure captions and system documentation were completed, followed by a final review of the report and project materials.",
        "All team members",
        "Completed",
    ],
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="8"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_font(run, size=9, bold=False, color="000000"):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.PORTRAIT
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.45)
section.right_margin = Inches(0.45)

doc.core_properties.title = "Weekly Progress Report"
doc.core_properties.subject = "Game engine project weekly development progress"

table = doc.add_table(rows=1, cols=len(HEADERS))
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.style = "Table Grid"

widths = [Inches(0.48), Inches(1.02), Inches(3.65), Inches(1.68), Inches(0.82)]

for index, text in enumerate(HEADERS):
    cell = table.rows[0].cells[index]
    cell.width = widths[index]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "1F4E78")
    set_cell_margins(cell, top=105, start=85, bottom=105, end=85)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    set_font(run, size=9, bold=True, color="FFFFFF")

set_repeat_table_header(table.rows[0])

for row_index, values in enumerate(ROWS, start=1):
    row = table.add_row()
    prevent_row_split(row)
    for col_index, value in enumerate(values):
        cell = row.cells[col_index]
        cell.width = widths[col_index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=90, start=85, bottom=90, end=85)
        if row_index % 2 == 0:
            set_cell_shading(cell, "EAF2F8")
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index in (2, 3) else WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(value)
        set_font(run, size=9, bold=False, color="000000")

set_table_borders(table)

# Fix the table grid so Word and LibreOffice use the intended column widths.
tbl_grid = table._tbl.tblGrid
for grid_col, width in zip(tbl_grid.gridCol_lst, widths):
    grid_col.set(qn("w:w"), str(int(width.inches * 1440)))

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
